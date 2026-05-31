from __future__ import annotations

from pathlib import Path


def read_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for .docx files") from exc
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def write_docx_text(source_path: Path, output_path: Path, text: str) -> None:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for .docx files") from exc
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document(source_path)
    paragraphs = document.paragraphs
    lines = text.splitlines() or [text]
    for index, paragraph in enumerate(paragraphs):
        paragraph.text = lines[index] if index < len(lines) else ""
    document.save(output_path)
