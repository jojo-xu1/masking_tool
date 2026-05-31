from __future__ import annotations

from pathlib import Path


def read_xlsx_blocks(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for .xlsx files") from exc
    workbook = load_workbook(path)
    values: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    values.append(cell.value)
    return values


def read_xlsx_text(path: Path) -> str:
    return "\n".join(read_xlsx_blocks(path))


def write_xlsx_blocks(source_path: Path, output_path: Path, blocks: list[str]) -> None:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for .xlsx files") from exc
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source_path)
    replacements = iter(blocks)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    cell.value = next(replacements, "")
    workbook.save(output_path)


def write_xlsx_text(source_path: Path, output_path: Path, text: str) -> None:
    write_xlsx_blocks(source_path, output_path, text.splitlines() or [text])
