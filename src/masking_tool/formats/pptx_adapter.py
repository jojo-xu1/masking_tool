from __future__ import annotations

from pathlib import Path


def read_pptx_blocks(path: Path) -> list[str]:
    try:
        from pptx import Presentation
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-pptx is required for .pptx files") from exc
    prs = Presentation(path)
    values: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text:
                            values.append(cell.text)
                continue
            if hasattr(shape, "text") and shape.text:
                values.append(shape.text)
    return values


def read_pptx_text(path: Path) -> str:
    return "\n".join(read_pptx_blocks(path))


def write_pptx_blocks(source_path: Path, output_path: Path, blocks: list[str]) -> None:
    try:
        from pptx import Presentation
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-pptx is required for .pptx files") from exc
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(source_path)
    replacements = iter(blocks)
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text:
                            cell.text = next(replacements, "")
                continue
            if hasattr(shape, "text") and shape.text:
                shape.text = next(replacements, "")
    prs.save(output_path)


def write_pptx_text(source_path: Path, output_path: Path, text: str) -> None:
    write_pptx_blocks(source_path, output_path, text.splitlines() or [text])
