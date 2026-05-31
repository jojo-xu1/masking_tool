from __future__ import annotations

from pathlib import Path

from masking_tool.core.models import DetectionResult, FileStatus, RiskLevel, TargetFile


REQUIRED_COLUMNS = [
    "No",
    "検出語句",
    "置換提案",
    "原文または前後の文脈",
    "情報カテゴリ",
    "リスクレベル",
    "判定理由",
    "推奨対応",
]
ALL_COLUMNS = REQUIRED_COLUMNS
RISK_FILLS = {
    "high": "F4CCCC",
    "medium": "FFF2CC",
    "low": "D9EAD3",
    "": "E7E6E6",
}
HEADER_FILL = "D9EAF7"
STATUS_FILL = "D9D9D9"
COLUMN_WIDTHS = {
    "A": 8,
    "B": 24,
    "C": 18,
    "D": 54,
    "E": 18,
    "F": 14,
    "G": 38,
    "H": 34,
}


def detection_row(detection: DetectionResult) -> list[str | int]:
    target = detection.target_file
    context = detection.context
    if target.relative_path:
        context = f"{context} [file: {target.relative_path}]"
    return [
        detection.no,
        detection.detected_text,
        detection.replacement,
        context,
        detection.category,
        detection.risk_level.value,
        detection.rule.judgment_reason,
        detection.rule.recommended_action,
    ]


def status_row(no: int, target: TargetFile) -> list[str | int]:
    reason = target.failure_reason or target.status.value
    return [
        no,
        "",
        "",
        f"{target.status.value}: {target.relative_path}",
        "STATUS",
        "",
        reason,
        _recommended_action_for_status(target),
    ]


def write_report(path: Path, detections: list[DetectionResult], targets: list[TargetFile]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to write Excel reports") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "検出結果"
    sheet.append(ALL_COLUMNS)
    sheet.freeze_panes = "A2"
    rows: list[list[str | int]] = [detection_row(d) for d in detections]
    no = len(rows) + 1
    detected_targets = {id(d.target_file) for d in detections}
    for target in targets:
        if id(target) not in detected_targets or target.status != FileStatus.PROCESSED or target.failure_reason:
            rows.append(status_row(no, target))
            no += 1
    for row in rows:
        sheet.append(row)
        risk = str(row[5])
        fill_color = STATUS_FILL if str(row[4]) == "STATUS" else RISK_FILLS.get(risk, RISK_FILLS[""])
        fill = PatternFill(fill_type="solid", fgColor=fill_color)
        for cell in sheet[sheet.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for column, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:H{max(sheet.max_row, 1)}"
    workbook.save(path)


def _recommended_action_for_status(target: TargetFile) -> str:
    if target.status == FileStatus.SKIPPED_UNSUPPORTED:
        return "Review file type if masking is required"
    if target.status == FileStatus.SKIPPED_OUT_OF_SCOPE:
        return "Use supported text-based content only"
    if target.status == FileStatus.FAILED:
        return "Review failure reason and retry"
    if target.status == FileStatus.NO_REPLACEMENT:
        return "No masking action required"
    return "Review processing status"
