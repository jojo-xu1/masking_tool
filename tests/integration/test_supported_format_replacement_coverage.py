from pathlib import Path
from importlib.resources import files

import pytest

from masking_tool.config.loader import load_rule_file, load_rule_files
from masking_tool.core.input_selection import make_folder_selection
from masking_tool.core.models import FileStatus
from masking_tool.core.processor import process


RULES = "tests/fixtures/rules/test_rules.yml"


def _default_rule_paths() -> list[Path]:
    defaults = files("masking_tool_defaults")
    return [Path(str(defaults / name)) for name in ("en.yml", "ja.yml", "zh.yml")]


def test_folder_replaces_log_and_csv_files(tmp_path: Path) -> None:
    source = tmp_path / "input"
    source.mkdir()
    (source / "app.log").write_text("INFO owner=Alice Smith email=alice@example.com", encoding="utf-8")
    (source / "contacts.csv").write_text("name,email\nAlice Smith,alice@example.com", encoding="utf-8")

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_file(RULES))

    statuses = {result.target.relative_path.as_posix(): result.target.status for result in results}
    assert statuses == {
        "app.log": FileStatus.PROCESSED,
        "contacts.csv": FileStatus.PROCESSED,
    }
    assert (output_dir / "files" / "app.log").read_text(encoding="utf-8") == "INFO owner=PERSON_001 email=EMAIL_001"
    assert (output_dir / "files" / "contacts.csv").read_text(encoding="utf-8") == "name,email\nPERSON_001,EMAIL_001"


def test_docx_replaces_paragraph_and_table_cell_text(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from masking_tool.formats.docx_adapter import read_docx_text

    source = tmp_path / "input"
    source.mkdir()
    document = docx.Document()
    document.add_paragraph("Contact Alice Smith at alice@example.com")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Approver Alice Smith uses alice@example.com"
    document.save(source / "office.docx")

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_file(RULES))

    assert results[0].target.status == FileStatus.PROCESSED
    masked_text = read_docx_text(output_dir / "files" / "office.docx")
    assert "Alice Smith" not in masked_text
    assert "alice@example.com" not in masked_text
    assert masked_text.count("PERSON_001") == 2
    assert masked_text.count("EMAIL_001") == 2


def test_xlsx_replaces_each_string_cell_without_losing_multiline_cell_text(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    from masking_tool.formats.xlsx_adapter import read_xlsx_text

    source = tmp_path / "input"
    source.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Contact Alice Smith\nalice@example.com"
    sheet["B1"] = "Backup Alice Smith alice@example.com"
    workbook.save(source / "book.xlsx")

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_file(RULES))

    assert results[0].target.status == FileStatus.PROCESSED
    masked_text = read_xlsx_text(output_dir / "files" / "book.xlsx")
    assert "Alice Smith" not in masked_text
    assert "alice@example.com" not in masked_text
    assert "Contact PERSON_001\nEMAIL_001" in masked_text
    assert "Backup PERSON_001 EMAIL_001" in masked_text


def test_pptx_replaces_textbox_and_table_cell_text(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    from masking_tool.formats.pptx_adapter import read_pptx_text

    source = tmp_path / "input"
    source.mkdir()
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Owner Alice Smith alice@example.com"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Reviewer Alice Smith alice@example.com"
    presentation.save(source / "slides.pptx")

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_file(RULES))

    assert results[0].target.status == FileStatus.PROCESSED
    masked_text = read_pptx_text(output_dir / "files" / "slides.pptx")
    assert "Alice Smith" not in masked_text
    assert "alice@example.com" not in masked_text
    assert masked_text.count("PERSON_001") == 2
    assert masked_text.count("EMAIL_001") == 2


def test_japanese_log_and_csv_samples_replace_person_names(tmp_path: Path) -> None:
    source = tmp_path / "input"
    source.mkdir()
    source.joinpath("ja_app.log").write_text(
        Path("docs/test-files/by-language/ja/ja_app.log").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    source.joinpath("ja_contacts.csv").write_text(
        Path("docs/test-files/by-language/ja/ja_contacts.csv").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_files(_default_rule_paths()))

    assert {result.target.relative_path.as_posix(): result.target.applied_language for result in results} == {
        "ja_app.log": "ja",
        "ja_contacts.csv": "ja",
    }
    for filename in ("ja_app.log", "ja_contacts.csv"):
        masked = (output_dir / "files" / filename).read_text(encoding="utf-8")
        assert "山田太郎" not in masked
        assert "佐藤花子" not in masked
        assert "PERSON_001" in masked
        assert "PERSON_002" in masked


def test_pdf_replacement_preserves_existing_pages_and_public_content(tmp_path: Path) -> None:
    fitz = pytest.importorskip("fitz")
    from masking_tool.formats.pdf_adapter import read_pdf_text

    source = tmp_path / "input"
    source.mkdir()
    pdf_path = source / "report.pdf"
    document = fitz.open()
    page1 = document.new_page()
    page1.insert_text((72, 72), "Public heading remains")
    page1.insert_text((72, 100), "Contact Alice Smith at alice@example.com")
    page2 = document.new_page()
    page2.insert_text((72, 72), "Second page public content remains")
    document.save(pdf_path)

    output_dir, results = process(make_folder_selection(source, tmp_path / "output"), load_rule_file(RULES))
    output_pdf = output_dir / "files" / "report.pdf"

    assert results[0].target.status == FileStatus.PROCESSED
    output_document = fitz.open(output_pdf)
    assert output_document.page_count == 2
    text = read_pdf_text(output_pdf)
    assert "Public heading remains" in text
    assert "Second page public content remains" in text
    assert "Alice Smith" not in text
    assert "alice@example.com" not in text
    assert "PERSON_001" in text
    assert "EMAIL_001" in text
