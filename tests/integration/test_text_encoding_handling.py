from pathlib import Path

import pytest

from masking_tool.core.input_selection import make_file_selection
from masking_tool.core.models import MaskingRule, RiskLevel, RuleType
from masking_tool.core.processor import process


def _explicit_rule() -> MaskingRule:
    return MaskingRule(
        "secret",
        "ja",
        True,
        RuleType.EXPLICIT,
        "SECRET",
        RiskLevel.HIGH,
        "explicit test secret",
        "replace",
        0,
        literal="秘密情報",
    )


def test_utf8_bom_is_not_in_detected_text_replacement_or_report_context(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    source = tmp_path / "bom.txt"
    source.write_text("秘密情報を確認しました。", encoding="utf-8-sig")

    output_dir, results = process(make_file_selection(source, "ja", tmp_path / "output"), [_explicit_rule()])

    assert results[0].detections[0].detected_text == "秘密情報"
    assert results[0].detections[0].replacement == "SECRET_001"
    assert "\ufeff" not in results[0].detections[0].context

    masked = (output_dir / "files" / "bom.txt").read_text(encoding="utf-8")
    assert masked == "SECRET_001を確認しました。"
    assert "\ufeff" not in masked

    workbook = load_workbook(output_dir / "機密情報検出結果.xlsx")
    row = next(workbook.active.iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[1] == "秘密情報"
    assert row[2] == "SECRET_001"
    assert "\ufeff" not in str(row[3])
