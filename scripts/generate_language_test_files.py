from __future__ import annotations

from pathlib import Path


ROOT = Path("docs/test-files/by-language")


SAMPLES = {
    "en": {
        "title": "English masking test memo",
        "person": "Alice Smith",
        "second_person": "John Miller",
        "email": "alice.smith@example.com",
        "phone": "202-555-0143",
        "mobile": "2025550143",
        "date": "2026-05-31",
        "postal": "10001",
        "address": "123 Market Street, New York, NY",
        "out_of_scope_postal": "100-0001",
        "project": "Project Falcon",
        "body": [
            "Owner: Alice Smith",
            "Reviewer: John Miller",
            "Contact email: alice.smith@example.com",
            "Primary phone: 202-555-0143",
            "Repeated phone: 202-555-0143",
            "Compact phone: 2025550143",
            "Review date: 2026-05-31",
            "Postal code: 10001",
            "Address: 123 Market Street, New York, NY",
            "Combined: Alice Smith phone 202-555-0143 ZIP 10001 Address: 123 Market Street, New York, NY",
            "Japan postal out of scope for English: 100-0001",
            "Ambiguous location note: New York office",
            "Sensitive project: Project Falcon",
            "Alice Smith approved the report.",
            "Alice Smith requested a second review.",
        ],
        "unsupported": "# Unsupported English sample\n\n- Alice Smith\n- alice.smith@example.com\n- 202-555-0143\n",
    },
    "ja": {
        "title": "日本語マスキングテストメモ",
        "person": "山田太郎",
        "second_person": "佐藤花子",
        "email": "yamada.taro@example.jp",
        "phone": "03-1234-5678",
        "mobile": "09012345678",
        "date": "2026-05-31",
        "postal": "100-0001",
        "address": "東京都千代田区千代田1-1",
        "out_of_scope_postal": "10001",
        "project": "極秘計画",
        "body": [
            "担当者: 山田太郎",
            "確認者: 佐藤花子",
            "連絡先メール: yamada.taro@example.jp",
            "代表電話: 03-1234-5678",
            "再掲電話: 03-1234-5678",
            "携帯電話: 09012345678",
            "確認日: 2026-05-31",
            "郵便番号: 100-0001",
            "住所: 東京都千代田区千代田1-1",
            "複合行: 山田太郎 電話 03-1234-5678 〒100-0001 住所: 東京都千代田区千代田1-1",
            "日本語では対象外の米国ZIP: 10001",
            "曖昧な場所メモ: 東京オフィス",
            "機密案件: 極秘計画",
            "山田太郎が草案を承認しました。",
            "山田太郎が再確認を依頼しました。",
        ],
        "unsupported": "# 未対応の日本語サンプル\n\n- 山田太郎\n- yamada.taro@example.jp\n- 03-1234-5678\n",
    },
    "zh": {
        "title": "中文脱敏测试备忘录",
        "person": "张伟",
        "second_person": "李雷",
        "email": "zhang.wei@example.cn",
        "phone": "+86 10 1234 5678",
        "mobile": "13800138000",
        "date": "2026-05-31",
        "postal": "100000",
        "address": "北京市朝阳区建国路88号",
        "out_of_scope_postal": "10001",
        "project": "秘密项目",
        "body": [
            "负责人: 张伟",
            "复核人: 李雷",
            "联系邮箱: zhang.wei@example.cn",
            "固定电话: +86 10 1234 5678",
            "重复电话: +86 10 1234 5678",
            "手机号码: 13800138000",
            "确认日期: 2026-05-31",
            "邮政编码: 100000",
            "地址: 北京市朝阳区建国路88号",
            "复合行: 张伟 电话 13800138000 邮政编码: 100000 地址: 北京市朝阳区建国路88号",
            "中文规则不处理美国ZIP: 10001",
            "模糊地点备注: 北京办公室",
            "敏感项目: 秘密项目",
            "张伟批准了草案。",
            "张伟要求再次复核。",
        ],
        "unsupported": "# 不支持的中文样本\n\n- 张伟\n- zhang.wei@example.cn\n- +86 10 1234 5678\n",
    },
}


def write_text_family(folder: Path, language: str, sample: dict[str, object]) -> None:
    body = list(sample["body"])
    (folder / f"{language}_sample.txt").write_text("\n".join([str(sample["title"]), "", *body]) + "\n", encoding="utf-8-sig")
    csv_rows = [
        "id,name,email,phone,postal,address,note,date",
        f"1,{sample['person']},{sample['email']},{sample['phone']},{sample['postal']},{sample['address']},{sample['project']},{sample['date']}",
        f"2,{sample['second_person']},{sample['email']},{sample['mobile']},{sample['out_of_scope_postal']},{sample['address']},{sample['project']},{sample['date']}",
    ]
    (folder / f"{language}_contacts.csv").write_text("\n".join(csv_rows) + "\n", encoding="utf-8-sig")
    log_rows = [
        f"INFO owner={sample['person']} phone={sample['phone']} date={sample['date']}",
        f"WARN reviewer={sample['second_person']} mobile={sample['mobile']} postal={sample['postal']} address={sample['address']}",
        f"INFO project={sample['project']} email={sample['email']}",
    ]
    (folder / f"{language}_app.log").write_text("\n".join(log_rows) + "\n", encoding="utf-8-sig")
    (folder / f"unsupported_{language}.md").write_text(str(sample["unsupported"]), encoding="utf-8-sig")


def write_docx(folder: Path, language: str, sample: dict[str, object]) -> None:
    from docx import Document

    document = Document()
    document.add_heading(str(sample["title"]), level=1)
    for line in sample["body"]:
        document.add_paragraph(str(line))
    document.save(folder / f"{language}_office.docx")


def write_xlsx(folder: Path, language: str, sample: dict[str, object]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sample"
    sheet.append(["field", "value"])
    for key in ("title", "person", "second_person", "email", "phone", "mobile", "date", "postal", "address", "project"):
        sheet.append([key, sample[key]])
    workbook.save(folder / f"{language}_workbook.xlsx")


def write_pptx(folder: Path, language: str, sample: dict[str, object]) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9.0), Inches(0.7))
    title_box.text_frame.text = str(sample["title"])
    title_box.text_frame.paragraphs[0].font.size = Pt(24)
    body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(9.0), Inches(5.5))
    frame = body_box.text_frame
    frame.clear()
    for index, line in enumerate(sample["body"][:8]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(line)
        paragraph.font.size = Pt(16)
    presentation.save(folder / f"{language}_presentation.pptx")


def write_pdf(folder: Path, language: str, sample: dict[str, object]) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page(width=595, height=842)
    fontname = "helv" if language == "en" else "china-s"
    text = "\n".join([str(sample["title"]), "", *[str(line) for line in sample["body"]]])
    page.insert_textbox((50, 50, 545, 792), text, fontsize=11, fontname=fontname)
    document.save(folder / f"{language}_text_pdf.pdf")


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    for language, sample in SAMPLES.items():
        folder = ROOT / language
        folder.mkdir(parents=True, exist_ok=True)
        write_text_family(folder, language, sample)
        write_docx(folder, language, sample)
        write_xlsx(folder, language, sample)
        write_pptx(folder, language, sample)
        write_pdf(folder, language, sample)


if __name__ == "__main__":
    main()
